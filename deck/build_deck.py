"""Build the Q-Hack India 2026 Round 1 deck for QSDTS on the official template.

    python deck/build_deck.py

Edits the organiser's template in place: every logo, mascot, cloud, footer,
section pill and the embedded IBM Plex Sans font are kept exactly. Only the
template's placeholder prompts are removed and replaced with content.

Every number comes from results/results.json, produced by
experiments/run_experiments.py. Every text box is checked against real
IBM Plex Sans metrics, and the build fails if any text would overflow.
"""

from __future__ import annotations

import json
import os
import sys
from pathlib import Path

from lxml import etree
from PIL import Image, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
RES = ROOT / "results"
# The organisers' template is not redistributed here: place it at deck/Q-Hack_India26.pptx
# or point QHACK_TEMPLATE at it. Fonts as in experiments/figures.py.
TEMPLATE = Path(os.environ.get("QHACK_TEMPLATE", ROOT / "deck" / "Q-Hack_India26.pptx"))
OUT = Path(os.environ.get("QHACK_OUT", ROOT.parent / "12-QHACK-ROUND1-QSDTS.pptx"))
FONT_DIR = Path(os.environ.get("PLEX_DIR", ROOT / "deck" / "fonts"))

FONT = "IBM Plex Sans"
PINK, PURPLE, INDIGO = "FA7CB5", "8041F9", "32145E"
RED, INK, MUTED, WHITE = "E5484D", "111111", "4A4A4A", "FFFFFF"
SOFT_PINK = "FDE3EF"

TEAM = "DoodleByte"
LEAD = "Gowtham K"
TRACK = "Quantum Security & Cryptography"
TITLE = "QSDTS: Keyless Quantum-Secure File Transfer"

R = json.loads(Path(os.environ.get("QHACK_RESULTS", RES / "results.json")).read_text())
PROBLEMS: list[str] = []

# --------------------------------------------------------------------------- fit check
_FONTS: dict[tuple[bool, float], ImageFont.FreeTypeFont] = {}


def _font(bold: bool, size: float) -> ImageFont.FreeTypeFont:
    key = (bold, size)
    if key not in _FONTS:
        name = "IBMPlexSans-Bold.ttf" if bold else "IBMPlexSans-Regular.ttf"
        # 10x oversampling: measure in tenths of a point
        _FONTS[key] = ImageFont.truetype(str(FONT_DIR / name), int(size * 10))
    return _FONTS[key]


def _lines_needed(text: str, size: float, bold: bool, width_pt: float) -> int:
    f = _font(bold, size)
    n = 0
    for para in text.split("\n"):
        words, line = para.split(" "), ""
        n += 1
        for w in words:
            trial = (line + " " + w).strip()
            # 4% headroom: renderers differ slightly in kerning and hinting
            if f.getlength(trial) / 10 > width_pt * 0.96 and line:
                n += 1
                line = w
            else:
                line = trial
    return n


# --------------------------------------------------------------------------- primitives
def _style_run(run, size, color=INK, bold=False, italic=False):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs", "a:sym"):
        el = rpr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rpr, qn(tag))
        el.set("typeface", FONT)


def _strip_style(shape):
    """Remove the theme style reference so no inherited outline or shadow appears."""
    st = shape._element.find(qn("p:style"))
    if st is not None:
        shape._element.remove(st)


def text(slide, x, y, w, h, paras, size=17, color=INK, align="l", anchor="t",
         spacing=1.18, gap=6, name="text"):
    """paras: list of str, or list of list[(text, {bold, color, italic, size})]."""
    if y > 2.9:  # everything below the section title must respect the zone
        check_zone(name, x, y, w, h)
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", 0)
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    total_lines = 0.0
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[align]
        p.line_spacing = spacing
        if i:
            p.space_before = Pt(gap)
        runs = [(para, {})] if isinstance(para, str) else para
        plain, psize, pbold = "", size, False
        for t, o in runs:
            r = p.add_run()
            r.text = t
            sz = o.get("size", size)
            _style_run(r, sz, o.get("color", color), o.get("bold", False), o.get("italic", False))
            plain += t
            psize = max(psize, sz)
            pbold = pbold or o.get("bold", False)
        lines = _lines_needed(plain, psize, pbold, w * 72)
        total_lines += lines * psize * spacing * 1.17 / 72 + (gap / 72 if i else 0)
    if total_lines > h + 0.02:
        PROBLEMS.append(f"{name}: needs {total_lines:.2f}in, box is {h:.2f}in")
    return tb


def box(slide, x, y, w, h, fill=WHITE, radius=0.12, line=None):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y),
                                 Inches(w), Inches(h))
    shp.adjustments[0] = min(0.5, radius / min(w, h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = RGBColor.from_string(fill)
    if line:
        shp.line.color.rgb = RGBColor.from_string(line)
        shp.line.width = Pt(1.5)
    else:
        shp.line.fill.background()
    _strip_style(shp)
    shp.shadow.inherit = False
    return shp


def pill(slide, x, y, label, size=30, fill=PINK, h=0.66, pad=0.32, max_w=None, lines=None):
    """A template-style pink pill sized to its text."""
    lines = lines or [label]
    f = _font(False, size)
    w = max(f.getlength(t) / 10 for t in lines) / 72 + 2 * pad
    if max_w and w > max_w:
        PROBLEMS.append(f"pill '{label[:30]}' is {w:.2f}in, limit {max_w:.2f}in")
    shp = box(slide, x, y, w, h, fill=fill, radius=h / 2)
    tf = shp.text_frame
    tf.word_wrap = False
    tf.margin_left = tf.margin_right = Inches(pad)
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, t in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.line_spacing = 1.0
        r = p.add_run()
        r.text = t
        _style_run(r, size)
    return shp


def dot(slide, x, y, d=0.17, fill=PINK):
    s = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    s.fill.solid()
    s.fill.fore_color.rgb = RGBColor.from_string(fill)
    s.line.fill.background()
    _strip_style(s)
    return s


def badge(slide, x, y, label, d=0.62, fill=PINK, color=INK, size=20):
    s = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    s.fill.solid()
    s.fill.fore_color.rgb = RGBColor.from_string(fill)
    s.line.fill.background()
    _strip_style(s)
    tf = s.text_frame
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", 0)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = label
    _style_run(r, size, color, bold=True)
    return s


def image(slide, path, x, y, w=None, h=None):
    with Image.open(path) as im:
        pw, ph = im.size
    if w and not h:
        h = w * ph / pw
    elif h and not w:
        w = h * pw / ph
    slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))
    return x, y, w, h


def keep_titles_on_one_line(slide):
    """The template's section titles sit in text boxes sized to the exact width of
    the text in Canva's embedded font. Any renderer with marginally wider metrics
    (Office's own cloud copy of IBM Plex Sans, for one) wraps them mid-word. With
    wrapping off, a centred title stays on one line, centred on its pill."""
    for shp in slide.shapes:
        if shp.has_text_frame and shp.text_frame.text.strip() and Emu(shp.top).inches < 2.9:
            shp.text_frame._txBody.find(qn("a:bodyPr")).set("wrap", "none")


def remove_prompt(slide):
    """Delete the template's question text (the only text box below the title
    that is not the footer). Works on both the grey and the purple variants,
    whose prompt boxes have different shape ids."""
    for shp in list(slide.shapes):
        if (shp.has_text_frame and Emu(shp.top).inches > 2.9
                and shp.text_frame.text.strip()
                and "Q-Hack India" not in shp.text_frame.text):
            shp._element.getparent().remove(shp._element)


def remove(slide, *ids):
    for shp in list(slide.shapes):
        if shp.shape_id in ids:
            shp._element.getparent().remove(shp._element)


def check_zone(name, x, y, w, h):
    """Content must stay clear of the logo strip, the footer, and the mascot."""
    if y < 2.95 or y + h > 10.45 or x < 0.7 or x + w > 19.4:
        PROBLEMS.append(f"{name}: outside the content zone ({x:.2f},{y:.2f},{w:.2f},{h:.2f})")
    if x + w > 17.75 and y + h > 9.2:
        PROBLEMS.append(f"{name}: collides with the bottom-right mascot")


def card(slide, x, y, w, h, fill=WHITE, name="card", **kw):
    check_zone(name, x, y, w, h)
    return box(slide, x, y, w, h, fill=fill, **kw)


# --------------------------------------------------------------------------- numbers
E0, E1, E2, E3, E4 = (R["e0_validation"], R["e1_efficiency"], R["e2_detection"],
                      R["e3_invariant"], R["e4_transfer"])
SHOW = R["showcase"]
QBER_ATTACK = 100 * E0["dl04_attack"]["qber"]
QBER_BITS = E0["dl04_attack"]["bits"]
ATTACK_QBER_SHOW = 100 * next(e["qber"] for e in SHOW["events"] if e["kind"] == "abort")
ATTACK_FRAME = 1 + next(e["frame"] for e in SHOW["events"] if e["kind"] == "abort")
RATIO = E1["dl04_per_qubit"] / E1["bb84_per_qubit"]
ADAPT, STATIC = E4["adaptive"], E4["static"]


def detect_at(frac):
    return 100 * next(r["detect_rate"] for r in E2["rows"] if abs(r["eve"] - frac) < 1e-9)


# --------------------------------------------------------------------------- slides
def slide1(s):
    remove(s, 17, 20, 23, 26, 35, 36, 37, 38)
    x, max_w = 0.54, 12.4
    pill(s, x, 4.30, f"Team Name: {TEAM}", max_w=max_w)
    pill(s, x, 5.22, f"Team Lead Name: {LEAD}", max_w=max_w)
    pill(s, x, 6.14, f"Track: {TRACK}", max_w=max_w)
    pill(s, x, 7.06, "Problem Statement title:", h=1.36, size=30, max_w=max_w,
         lines=["Problem Statement title:", TITLE])


def slide2(s):
    remove_prompt(s)
    text(s, 1.2, 3.1, 17.6, 1.05, [
        "Every secure channel today depends on a maths problem staying hard. "
        "Quantum computers break that assumption, and today's quantum fix still leaves a key to steal."
    ], size=22, align="c", name="s2 headline")
    cards = [
        ("Security rests on slowness",
         "RSA and elliptic-curve key exchange are safe only because factoring and discrete "
         "logarithms are slow on classical computers. Shor's algorithm solves both in "
         "polynomial time on a large quantum computer [4]."),
        ("Harvest now, decrypt later",
         "Encrypted traffic recorded today can be decrypted once such a machine exists. "
         "Health, genomic, government and financial data that must stay secret for decades "
         "is already exposed."),
        ("QKD still leaves a key",
         "Quantum key distribution secures key agreement with physics, but the message still "
         "travels classically, and keys must be stored and managed at every node: a "
         "high-value target."),
    ]
    for i, (title, body) in enumerate(cards):
        x = 1.2 + i * 6.0
        card(s, x, 4.35, 5.6, 3.95, name=f"s2 card {i}")
        badge(s, x + 0.35, 4.62, str(i + 1))
        text(s, x + 1.15, 4.66, 4.2, 0.6, [title], size=20, name=f"s2 c{i} title",
             color=INDIGO)
        text(s, x + 0.35, 5.45, 4.95, 2.7, [body], size=17.5, color=MUTED,
             name=f"s2 c{i} body")
    card(s, 1.2, 8.6, 16.3, 1.2, fill=PINK, name="s2 gap")
    text(s, 1.55, 8.67, 15.6, 1.06, [[
        ("The gap: ", {"bold": True}),
        ("QSDC removes the key entirely, but it exists only in optics labs. "
         "A 2025 review of 11 quantum-network simulators found none that implement it [9].", {}),
    ]], size=18.5, anchor="m", name="s2 gap text")


def slide3(s):
    remove_prompt(s)
    text(s, 1.2, 3.15, 10.4, 0.6, ["Who faces this problem"], size=23, color=INDIGO,
         name="s3 heading")
    rows = [
        ("Hospitals and genomics labs",
         "Records must stay confidential for a lifetime, longer than today's encryption is "
         "expected to last."),
        ("Government, defence and courts",
         "Archives that are valuable enough to harvest now and read later."),
        ("Banks and critical infrastructure",
         "Key stores at every network node are single, high-value targets."),
        ("Researchers and students",
         "QSDC networks cannot be built, compared or reproduced without a photonics lab, "
         "so the field stays closed."),
    ]
    y = 3.95
    for t, b in rows:
        card(s, 1.2, y, 10.4, 1.38, name=f"s3 row {t}")
        dot(s, 1.5, y + 0.3, d=0.24)
        text(s, 1.95, y + 0.17, 9.4, 0.5, [t], size=19, name=f"s3 {t}")
        text(s, 1.95, y + 0.66, 9.4, 0.66, [b], size=15.5, color=MUTED, name=f"s3 {t} body")
        y += 1.55
    stats = [
        ("0", "keys created, stored or managed in QSDC: nothing to steal later"),
        ("11 → 0", "quantum-network simulators reviewed in 2025 that implement QSDC [9]"),
        ("2,000 km", "secure quantum communication targeted by India's National Quantum "
                     "Mission (2023) [12]"),
    ]
    y = 3.15
    for big, small in stats:
        card(s, 12.2, y, 5.5, 1.95, name=f"s3 stat {big}")
        text(s, 12.5, y + 0.12, 4.9, 0.85, [big], size=40, color=PURPLE, name=f"s3 big {big}")
        text(s, 12.5, y + 1.0, 4.9, 0.85, [small], size=14.5, color=MUTED,
             name=f"s3 small {big}")
        y += 2.1


def slide4(s):
    remove_prompt(s)
    text(s, 1.2, 3.05, 17.6, 1.0, [
        "QSDTS is a network of software quantum nodes that sends a file as quantum states, "
        "catches an eavesdropper before anything is encoded, and routes around them."
    ], size=21.5, align="c", name="s4 headline")
    card(s, 1.0, 4.25, 12.55, 4.75, name="s4 diagram card")
    x, y, w, h = image(s, RES / "diagram_dl04.png", 1.12, 4.38, w=12.3)
    check_zone("s4 diagram", x, y, w, h)
    chips = [
        (INDIGO, WHITE, "No key, ever", "The message itself is the quantum state (DL04 QSDC)."),
        (PINK, INK, "Caught before encoding", "A failed check means the message never touches the channel."),
        (WHITE, INDIGO, "Routed around", "A tapped link is dropped; the file still arrives, verified."),
    ]
    cy = 4.35
    for fill, col, head, body in chips:
        card(s, 13.85, cy, 3.85, 1.42, fill=fill, name=f"s4 chip {head}")
        text(s, 14.12, cy + 0.12, 3.35, 0.5, [head], size=18, color=col, name=f"s4 {head}")
        text(s, 14.12, cy + 0.58, 3.35, 0.78, [body], size=13.5, color=col,
             name=f"s4 {head} body")
        cy += 1.58
    text(s, 1.2, 9.35, 16.2, 0.9, [
        "Each hop is an independent DL04 session through a trusted relay, the model every "
        "deployed quantum network uses today. The same channel also runs BB84, so the two "
        "protocols are compared under identical conditions."
    ], size=15, color=INK, align="c", name="s4 footer")


def slide5(s):
    remove_prompt(s)
    rows, cols = 5, 5
    tx, ty, tw, th = 1.2, 3.2, 11.0, 4.6
    check_zone("s5 table", tx, ty, tw, th)
    tbl = s.shapes.add_table(rows, cols, Inches(tx), Inches(ty), Inches(tw), Inches(th)).table
    data = [
        ["", "RSA / ECC", "Post-quantum", "QKD", "QSDTS (QSDC)"],
        ["Security rests on", "maths hardness", "new, unproven maths hardness", "physics", "physics"],
        ["Broken by a better algorithm?", "yes (Shor)", "not known", "no", "no"],
        ["A key to store and steal?", "yes", "yes", "yes", "no key exists"],
        ["Detects eavesdropping?", "no", "no", "yes", "yes, before sending"],
    ]
    widths = [2.75, 1.9, 2.3, 1.55, 2.5]
    for c, wd in enumerate(widths):
        tbl.columns[c].width = Inches(wd)
    for r in range(rows):
        tbl.rows[r].height = Inches(0.92)
        for c in range(cols):
            cell = tbl.cell(r, c)
            cell.margin_left = cell.margin_right = Inches(0.12)
            cell.margin_top = cell.margin_bottom = Inches(0.04)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            fill = WHITE if c < 4 else SOFT_PINK
            if r == 0:
                fill = INDIGO if c < 4 else PINK
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor.from_string(fill)
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT if c == 0 else PP_ALIGN.CENTER
            run = p.add_run()
            run.text = data[r][c]
            head_col = WHITE if (r == 0 and c < 4) else INK
            _style_run(run, 15.5 if r else 16.5, head_col, bold=(r == 0 or c == 0 or c == 4))
            need = _lines_needed(data[r][c], 16.5, True, (widths[c] - 0.24) * 72)
            if need * 16.5 * 1.25 / 72 > 0.84:
                PROBLEMS.append(f"s5 cell {r},{c} wraps to {need} lines")
    # table style: drop the banding the default style adds
    tblpr = tbl._tbl.tblPr
    tblpr.set("firstRow", "0")
    tblpr.set("bandRow", "0")

    text(s, 1.2, 8.15, 11.0, 2.0, [[
        ("A classical bit can be copied perfectly and silently. A qubit cannot. ",
         {"italic": True, "color": INDIGO}),
        ("That single fact is the whole security argument, and it is why this problem "
         "needs quantum rather than better classical maths.", {"italic": True}),
    ]], size=19, name="s5 statement")

    card(s, 12.75, 3.2, 4.95, 5.85, name="s5 why card")
    text(s, 13.1, 3.42, 4.4, 0.6, ["Why only quantum works here"], size=19, color=INDIGO,
         name="s5 why head")
    items = [
        ("No-cloning", "An unknown qubit cannot be copied [5], so Eve cannot keep a copy "
                       "and forward the original unnoticed."),
        ("Measurement disturbs", "Reading a qubit in the wrong basis randomises it, leaving "
                                 "errors the receiver can count."),
        ("The 25% signature", "Eve picks the wrong basis half the time; then the bit reads "
                              "wrong half the time. 0.5 × 0.5 = 25% error."),
    ]
    y = 4.2
    for head, body in items:
        dot(s, 13.15, y + 0.14, d=0.22, fill=PURPLE)
        text(s, 13.55, y, 3.95, 0.5, [head], size=18, name=f"s5 {head}")
        text(s, 13.55, y + 0.48, 3.95, 1.12, [body], size=14, color=MUTED,
             name=f"s5 {head} body")
        y += 1.6


def slide6(s):
    remove_prompt(s)
    x, y, w, h = image(s, RES / "diagram_architecture.png", 1.2, 3.25, w=9.4)
    check_zone("s6 diagram", x, y, w, h)
    card(s, 11.0, 3.2, 7.6, 5.9, name="s6 card")
    text(s, 11.35, 3.4, 7.0, 0.6, ["How the parts work together"], size=20.5, color=INDIGO,
         name="s6 head")
    items = [
        ("Quantum. ", "Every qubit is a wire in a Qiskit circuit, simulated exactly on "
                      "Qiskit Aer. Fibre noise is a depolarizing channel; the eavesdropper "
                      "is a mid-circuit measurement, so the errors she causes come from "
                      "quantum mechanics, not from a formula."),
        ("Check before encode. ", "The quantum state is saved after Alice's check and "
                                  "restored only if the classical decision is PROCEED. "
                                  "An aborted session never creates a single encoding gate."),
        ("Classical. ", "NetworkX routing with link cost  distance ÷ (1−QBER/11%)², "
                        "so traffic leaves a degrading link early. Hamming(7,4) + CRC-32 per "
                        "hop, SHA-256 end to end."),
    ]
    yy = 4.15
    for head, body in items:
        dot(s, 11.4, yy + 0.12, d=0.2, fill=PINK)
        text(s, 11.8, yy, 6.55, 1.55, [[(head, {"bold": True, "color": INDIGO}), (body, {})]],
             size=15, name=f"s6 {head}")
        yy += 1.6
    card(s, 1.2, 9.3, 16.3, 0.85, fill=INDIGO, name="s6 stack")
    text(s, 1.5, 9.3, 15.7, 0.85, [[
        ("Stack   ", {"bold": True, "color": PINK}),
        ("Python 3.12  ·  Qiskit 2.5  ·  Qiskit Aer 0.17  ·  NetworkX  ·  NumPy  ·  "
         "Matplotlib  ·  pytest", {"color": WHITE}),
    ]], size=16.5, anchor="m", align="c", name="s6 stack text")


def slide7(s):
    remove_prompt(s)
    card(s, 1.2, 3.1, 16.4, 2.95, name="s7 circuit card")
    text(s, 1.5, 3.22, 15.8, 0.42, [
        "The real circuit every payload qubit runs through (qsdts/dl04.py): Bob prepares, "
        "Alice checks a subset, then encodes a 1 with Y on the qubits that remain"
    ], size=14, color=MUTED, align="c", name="s7 caption")
    x, y, w, h = image(s, RES / "diagram_circuit.png", 2.4, 3.7, w=14.0)
    check_zone("s7 circuit", x, y, w, h)

    top = 6.25
    card(s, 1.2, top, 7.55, 4.1, name="s7 list card")
    text(s, 1.5, top + 0.18, 7.0, 0.5, ["What runs today"], size=19.5, color=INDIGO,
         name="s7 list head")
    items = [
        "DL04 QSDC and BB84 on one shared quantum channel",
        "Intercept-resend eavesdropper at any strength, 0 to 100%",
        "Abort-before-encode, enforced by an automated test",
        "5-node network, trusted relays, adaptive rerouting",
        "File transfer: per-hop CRC, end-to-end SHA-256",
        "40 tests passing; every chart rebuilt by one script",
    ]
    yy = top + 0.82
    for it in items:
        dot(s, 1.55, yy + 0.11, d=0.17)
        text(s, 1.9, yy, 6.7, 0.42, [it], size=15, name=f"s7 item {it[:12]}")
        yy += 0.5

    card(s, 9.05, top, 8.55, 4.1, name="s7 network card")
    nx_, ny, nw, nh = image(s, RES / "diagram_network.png", 10.4, top + 0.06, w=5.85)
    check_zone("s7 network", nx_, ny, nw, nh)
    text(s, 9.3, top + 3.4, 8.05, 0.66, [
        f"Live run: Eve taps R1-Bob at frame {ATTACK_FRAME} of {SHOW['frames']}. QBER "
        f"{ATTACK_QBER_SHOW:.1f}%, aborted before encoding; R1 reroutes via R3; all "
        f"{SHOW['frames']} frames delivered, SHA-256 verified."
    ], size=13, color=MUTED, align="c", name="s7 network caption")


def slide8(s):
    remove_prompt(s)
    stats = [
        (f"{QBER_ATTACK:.1f}%", f"measured QBER under intercept-resend; theory says 25% "
                                f"({QBER_BITS:,} check bits)"),
        (f"{E3['message_bits_exposed']} bits",
         f"of message ever sent on a link that failed its check "
         f"({E3['aborted_sessions_and_transfers']} aborted runs)"),
        (f"{ADAPT['verified']}/{ADAPT['n']} vs {STATIC['verified']}/{STATIC['n']}",
         "files delivered and hash-verified under a mid-transfer attack, controller on vs off"),
    ]
    for i, (big, small) in enumerate(stats):
        x = 1.2 + i * 5.55
        card(s, x, 3.15, 5.25, 1.75, name=f"s8 stat {i}")
        text(s, x + 0.3, 3.22, 4.7, 0.78, [big], size=34, color=PURPLE, name=f"s8 big {i}")
        text(s, x + 0.3, 4.02, 4.7, 0.8, [small], size=13.5, color=MUTED, name=f"s8 small {i}")

    card(s, 1.2, 5.1, 8.55, 5.2, name="s8 qber card")
    x, y, w, h = image(s, RES / "chart_qber_attack.png", 1.35, 5.2, w=8.0)
    check_zone("s8 chart qber", x, y, w, h)
    text(s, 1.45, y + h + 0.02, 8.05, 0.8, [
        f"Measured QBER tracks theory. Attacks on 50% or more of qubits are stopped on "
        f"{detect_at(0.5):.0f}% of blocks; at 20% only {detect_at(0.2):.0f}% are, because a weak "
        f"attacker hides in fibre noise within one block. That is our Round 2 question."
    ], size=12, color=MUTED, name="s8 cap qber")

    card(s, 10.0, 5.1, 7.55, 5.2, name="s8 eff card")
    x2, y2, w2, h2 = image(s, RES / "chart_efficiency.png", 10.45, 5.25, w=6.65)
    check_zone("s8 chart eff", x2, y2, w2, h2)
    text(s, 10.25, y2 + h2 + 0.05, 7.05, 0.85, [
        f"{RATIO:.1f}x more message bits per qubit prepared, but about equal per fibre "
        f"transmission, because DL04 sends each qubit twice. The win is that no key exists, "
        f"not raw throughput."
    ], size=12.5, color=MUTED, name="s8 cap eff")


def slide9(s):
    remove_prompt(s)
    steps = [
        ("Hardware-grounded",
         "Calibrate the channel model on IBM Quantum hardware (T1, T2, gate and readout "
         "errors) and run the core DL04 circuits on a real device.",
         "Results tied to a real device's error rates, not chosen parameters."),
        ("Stronger adversaries",
         "Partial and adaptive eavesdroppers, detection across many blocks instead of one, "
         "photon-number-splitting attacks and decoy states.",
         "The weakest attacker the network reliably catches, and how fast."),
        ("Scale and baselines",
         "E91 with a CHSH test as a third protocol; 10 to 50 node topologies; per-hop "
         "latency and throughput under load.",
         "QSDC vs BB84 vs E91 under one common noise model."),
        ("Live demo and release",
         "A web dashboard to tap any link and watch QBER, policy and route react in real "
         "time; documented open-source release.",
         "A live demo judges can drive themselves, end to end."),
    ]
    line = s.shapes.add_connector(1, Inches(1.8), Inches(3.62), Inches(17.2), Inches(3.62))
    line.line.color.rgb = RGBColor.from_string(PINK)
    line.line.width = Pt(4)
    for i, (head, body, outcome) in enumerate(steps):
        x = 1.2 + i * 4.2
        badge(s, x + 1.65, 3.3, str(i + 1), d=0.66, fill=PINK, size=21)
        card(s, x, 4.2, 3.95, 4.05, name=f"s9 step {i}")
        text(s, x + 0.3, 4.4, 3.4, 0.55, [head], size=19, color=INDIGO, name=f"s9 head {i}")
        text(s, x + 0.3, 5.0, 3.4, 1.75, [body], size=15, color=MUTED, name=f"s9 body {i}")
        box(s, x + 0.2, 6.85, 3.55, 1.22, fill=SOFT_PINK, radius=0.1)
        text(s, x + 0.38, 6.93, 3.2, 1.08, [[("Outcome  ", {"bold": True, "color": INDIGO}),
                                            (outcome, {})]],
             size=14, anchor="m", name=f"s9 outcome {i}")
    card(s, 1.2, 8.55, 16.3, 1.55, fill=INDIGO, name="s9 limits")
    text(s, 1.55, 8.6, 15.6, 1.45, [[
        ("Limits we will keep stating.  ", {"bold": True, "color": PINK}),
        ("Relays are trusted and see the plaintext, as in every deployed quantum network "
         "today; removing that needs quantum repeaters. All results so far are simulation "
         "on Qiskit Aer until step 1 lands.", {"color": WHITE}),
    ]], size=16, anchor="m", name="s9 limits text")


REFS = [
    'F.-G. Deng and G. L. Long, "Secure direct communication with a quantum one-time pad," Phys. Rev. A, vol. 69, 052319, 2004.',
    'G. L. Long and X. S. Liu, "Theoretically efficient high-capacity quantum-key-distribution scheme," Phys. Rev. A, vol. 65, 032302, 2002.',
    'C. H. Bennett and G. Brassard, "Quantum cryptography: Public key distribution and coin tossing," Proc. IEEE ICCSSP, Bangalore, 1984, pp. 175-179.',
    'P. W. Shor, "Algorithms for quantum computation: discrete logarithms and factoring," Proc. 35th FOCS, 1994, pp. 124-134.',
    'W. K. Wootters and W. H. Zurek, "A single quantum cannot be cloned," Nature, vol. 299, pp. 802-803, 1982.',
    'Z. Qi et al., "A 15-user quantum secure direct communication network," Light: Sci. Appl., vol. 10, 183, 2021.',
    'M. Wang et al., "Experimental demonstration of secure relay in quantum secure direct communication network," Entropy, vol. 25, 1548, 2023.',
    'S. Zhang and C. Zheng, "Quantum secure direct communication technology-enhanced time-sensitive networks," Entropy, vol. 27, 221, 2025.',
    'R. J. Hayek, J. Chung and R. Kettimuthu, "A review of software for designing and operating quantum networks," arXiv:2510.00203, 2025.',
    'M. Sasaki et al., "Field test of quantum key distribution in the Tokyo QKD Network," Opt. Express, vol. 19, pp. 10387-10409, 2011.',
    "Qiskit and Qiskit Aer: github.com/Qiskit/qiskit, github.com/Qiskit/qiskit-aer. NetworkX: networkx.org.",
    'Press Information Bureau, Govt. of India, "Cabinet approves National Quantum Mission to scale-up scientific and industrial R&D for quantum technologies," 19 April 2023.',
]


def slide10(s):
    remove_prompt(s)
    half = 6
    for col in range(2):
        x = 1.2 + col * 8.45
        card(s, x, 3.1, 8.15, 5.35, name=f"s10 col {col}")
        paras = [[(f"[{i + 1}]  ", {"bold": True, "color": PURPLE}), (REFS[i], {})]
                 for i in range(col * half, (col + 1) * half)]
        text(s, x + 0.3, 3.3, 7.6, 5.0, paras, size=14, gap=8, color=INK,
             name=f"s10 refs {col}")
    card(s, 1.2, 8.65, 16.3, 1.5, fill=PINK, name="s10 ack")
    text(s, 1.5, 8.7, 15.7, 1.4, [[
        ("Acknowledgement.  ", {"bold": True}),
        ("QSDTS implements published protocols (DL04 [1], BB84 [3]) and a published idea, "
         "eavesdropper-triggered rerouting, demonstrated on the Tokyo QKD Network in 2010 "
         "[10]. The implementation, network layer, experiments and figures are our own work; "
         "no repository was copied or forked. Code, tests and a one-click Colab demo: "
         "github.com/gowtham472/qsdts. Built with AI assistance (Claude Code); design "
         "decisions, review and validation by the team.",
         {}),
    ]], size=14.5, anchor="m", name="s10 ack text")


def main() -> int:
    prs = Presentation(str(TEMPLATE))
    builders = [slide1, slide2, slide3, slide4, slide5, slide6, slide7, slide8, slide9, slide10]
    for i, (s, build) in enumerate(zip(prs.slides, builders, strict=True)):
        if i:
            keep_titles_on_one_line(s)
        build(s)
    if PROBLEMS:
        print("LAYOUT PROBLEMS:")
        for p in PROBLEMS:
            print("  -", p)
        return 1
    prs.save(str(OUT))
    print("wrote", OUT)
    return 0


if __name__ == "__main__":
    sys.exit(main())
